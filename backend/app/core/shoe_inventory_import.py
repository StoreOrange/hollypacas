"""Initial legacy inventory loading, only used by the bdzapatos routes."""
import hashlib
from decimal import Decimal, InvalidOperation
from io import BytesIO
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from sqlalchemy import func, text as sql_text
from ..models.inventory import (Producto, Linea, ColorCatalog, ShoeProductVariant,
    ShoeVariantStock, SaldoProducto, IngresoTipo, IngresoInventario, IngresoItem)

HEADERS = ['CODPRODUCTO', 'DESCRIPCION', 'MARCA', 'TAMANO', 'COLOR', 'LINEA',
           'CODBARRA', 'PRECIO', '001 - PRINCIPAL', '002 - KGF', '003 - KG']
WAREHOUSES = {'001 - PRINCIPAL': 'central', '002 - KGF': 'kgf', '003 - KG': 'kg'}

def text(value):
    return str(value).strip() if value is not None else ''

def number(value, row, column):
    raw = text(value)
    if raw.upper() in {'', 'NULL'}:
        return Decimal('0')
    try:
        result = Decimal(raw.replace(',', ''))
    except InvalidOperation:
        raise ValueError(f'Fila {row}: {column} debe ser numerico')
    if not result.is_finite() or result < 0 or result >= Decimal('10000000000'):
        raise ValueError(f'Fila {row}: {column} fuera de rango')
    return result

def parse(content):
    try:
        book = load_workbook(BytesIO(content), data_only=True)
    except Exception:
        raise ValueError('No se pudo leer el archivo Excel')
    sheet = book.active
    columns = [text(c.value).upper() for c in sheet[1]]
    if len(columns) != len(set(columns)):
        raise ValueError('Hay encabezados duplicados')
    missing = set(HEADERS) - set(columns)
    if missing:
        raise ValueError('Faltan columnas: ' + ', '.join(sorted(missing)))
    rows, seen = [], set()
    limits = {'CODPRODUCTO':60,'DESCRIPCION':200,'MARCA':80,'TAMANO':20,'COLOR':120,'LINEA':120,'CODBARRA':120}
    for index, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        if not any(text(v) for v in values):
            continue
        data = dict(zip(columns, values))
        row = {key: text(data.get(key)) for key in HEADERS}
        for key, limit in limits.items():
            if key != 'CODBARRA' and not row[key]:
                raise ValueError(f'Fila {index}: {key} es obligatorio')
            if len(row[key]) > limit:
                raise ValueError(f'Fila {index}: {key} supera {limit} caracteres')
        if row['CODPRODUCTO'] in seen:
            raise ValueError(f'Fila {index}: CODPRODUCTO duplicado')
        seen.add(row['CODPRODUCTO'])
        for key in ['PRECIO', *WAREHOUSES]:
            row[key] = number(data.get(key), index, key)
            if key in WAREHOUSES and row[key] != row[key].to_integral_value():
                raise ValueError(f'Fila {index}: {key} debe ser cantidad entera')
            if key == 'PRECIO' and row[key] != row[key].quantize(Decimal('.01')):
                raise ValueError(f'Fila {index}: PRECIO admite dos decimales')
        row['COSTO'] = None if text(data.get('COSTO')).upper() in {'','NULL'} else number(data['COSTO'],index,'COSTO')
        if row['COSTO'] is not None and row['COSTO'] != row['COSTO'].quantize(Decimal('.01')):
            raise ValueError(f'Fila {index}: COSTO admite dos decimales')
        rows.append(row)
    if not rows:
        raise ValueError('El archivo no contiene articulos')
    return rows

def template():
    book = Workbook()
    sheet = book.active
    sheet.title = 'Inventario'
    sheet.append(HEADERS + ['COSTO'])
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = 'A1:L1'
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        sheet.column_dimensions[cell.column_letter].width = 24
    for col in ['A','D','G']:
        for row in range(2, 1002):
            sheet[f'{col}{row}'].number_format = '@'
    notes = book.create_sheet('Instrucciones')
    for note in [
        'Copiar las 11 columnas originales en el mismo orden. COSTO es opcional y en cordobas.',
        'PRECIO en cordobas. NULL o celda vacia en existencias significa cero.',
        '001 - PRINCIPAL = Central; 002 - KGF = KGF; 003 - KG = KG. Esteli no se modifica.',
        'Cada CODPRODUCTO conserva su producto, talla, color y precio. CODBARRA es referencia compartida.',
        'Carga inicial: solo crea saldos de variantes nuevas. Una variante existente debe tener los mismos saldos.',
        'Repetir el mismo archivo no suma existencias. Para ajustes posteriores use movimientos de inventario.',
        'COSTO vacio conserva el existente; en articulos nuevos queda cero. No se calcula desde PRECIO.',
        'Todos los encabezados originales son obligatorios. CODBARRA puede quedar vacio.',
    ]:
        notes.append([note])
    notes.column_dimensions['A'].width = 140
    stream = BytesIO(); book.save(stream); stream.seek(0)
    return stream

def _insert(db, model, values, returning=None):
    result = []
    for start in range(0, len(values), 500):
        batch = values[start:start + 500]
        if returning and db.get_bind().dialect.name == 'postgresql':
            statement = model.__table__.insert().values(batch).returning(*[model.__table__.c[name] for name in returning])
            result.extend(db.execute(statement).all())
        elif returning:
            objects = [model(**row) for row in batch]
            db.add_all(objects)
            db.flush()
            result.extend(tuple(getattr(obj, name) for name in returning) for obj in objects)
        else:
            db.execute(model.__table__.insert(), batch)
    return result


def apply(db, rows, warehouses, rate, day, username):
    # SQL echo on large uploads overwhelms local disk and terminal output.
    db.get_bind().echo = False
    if db.get_bind().dialect.name == 'postgresql':
        locked = db.execute(sql_text('SELECT pg_try_advisory_xact_lock(802001)')).scalar()
        if not locked:
            raise ValueError('Otra importacion esta en curso. Espere a que termine antes de subir nuevamente')
    stores = {}
    for header, code in WAREHOUSES.items():
        matches = [w for w in warehouses if text(w.code).lower() == code or text(w.name).lower() == code]
        if len(matches) != 1:
            raise ValueError(f'Bodega {code}: no encontrada o ambigua')
        stores[header] = matches[0]
    codes = [row['CODPRODUCTO'] for row in rows]
    products, variants = {}, {}
    for start in range(0, len(codes), 1000):
        batch = codes[start:start + 1000]
        products.update((p.cod_producto, p) for p in db.query(Producto).filter(Producto.cod_producto.in_(batch)))
        variants.update((v.cod_variante, v) for v in db.query(ShoeProductVariant).filter(ShoeProductVariant.cod_variante.in_(batch)))
    colors = {c.nombre.lower(): c for c in db.query(ColorCatalog)}
    color_ids = {c.id: c for c in colors.values()}
    lines = {line.linea.lower(): line for line in db.query(Linea)}
    stocks = {}
    ids = [v.id for v in variants.values()]
    for start in range(0, len(ids), 1000):
        for stock in db.query(ShoeVariantStock).filter(ShoeVariantStock.variante_id.in_(ids[start:start + 1000])):
            stocks[(stock.variante_id, stock.bodega_id)] = Decimal(str(stock.existencia or 0))
    for row in rows:
        code = row['CODPRODUCTO']; variant = variants.get(code); product = products.get(code)
        if variant:
            if not product or variant.producto_id != product.id:
                raise ValueError(f'{code}: variante existente pertenece a otro producto; requiere conciliacion')
            if variant.talla != row['TAMANO'] or color_ids[variant.color_id].nombre.lower() != row['COLOR'].lower():
                raise ValueError(f'{code}: talla/color no coinciden con la variante existente')
            for header, store in stores.items():
                current = stocks.get((variant.id, store.id), Decimal('0'))
                if current != row[header]:
                    raise ValueError(f'{code}: saldo existente en {store.name} ({current}) distinto del archivo ({row[header]}). Use ajustes de inventario')
        elif product:
            raise ValueError(f'{code}: producto existente sin variante coincidente; requiere conciliacion')
    for row in rows:
        key = row['LINEA'].lower()
        if key not in lines:
            line = Linea(cod_linea='IMP-' + hashlib.sha256(key.encode()).hexdigest()[:40], linea=row['LINEA'], activo=True)
            db.add(line); lines[key] = line
        key = row['COLOR'].lower()
        if key not in colors:
            color = ColorCatalog(nombre=row['COLOR'], abreviatura=row['COLOR'][:20], activo=True)
            db.add(color); colors[key] = color
    db.flush()
    new_rows, product_values = [], []
    for row in rows:
        values = dict(descripcion=row['DESCRIPCION'], marca=row['MARCA'], linea_id=lines[row['LINEA'].lower()].id,
                      referencia_producto=row['CODBARRA'], precio_venta1=row['PRECIO'],
                      precio_venta1_usd=row['PRECIO']/rate if rate > 0 else None, tasa_cambio=rate if rate > 0 else None)
        if row['COSTO'] is not None: values['costo_producto'] = row['COSTO']
        product = products.get(row['CODPRODUCTO'])
        if product:
            for key, value in values.items(): setattr(product, key, value)
        else:
            values.setdefault('costo_producto', Decimal('0'))
            values.update(cod_producto=row['CODPRODUCTO'], activo=True)
            product_values.append(values); new_rows.append(row)
    db.flush()
    product_ids = dict((code, ident) for ident, code in _insert(db, Producto, product_values, ['id', 'cod_producto']))
    variant_values = [dict(producto_id=product_ids[row['CODPRODUCTO']], color_id=colors[row['COLOR'].lower()].id,
                           talla=row['TAMANO'], cod_variante=row['CODPRODUCTO'], activo=True) for row in new_rows]
    variant_ids = dict((code, ident) for ident, code in _insert(db, ShoeProductVariant, variant_values, ['id', 'cod_variante']))
    saldo_values, stock_values, incoming = [], [], {}
    for row in new_rows:
        code = row['CODPRODUCTO']; pid = product_ids[code]; vid = variant_ids[code]
        saldo_values.append(dict(producto_id=pid, existencia=sum(row[h] for h in stores)))
        for header, store in stores.items():
            qty = row[header]
            stock_values.append(dict(variante_id=vid, bodega_id=store.id, existencia=qty))
            if qty: incoming.setdefault(store.id, []).append((pid, vid, qty, row['COSTO'] or Decimal('0')))
    _insert(db, SaldoProducto, saldo_values)
    _insert(db, ShoeVariantStock, stock_values)
    if incoming:
        kind = db.query(IngresoTipo).filter(func.lower(IngresoTipo.nombre).like('%apertura%')).first()
        if not kind:
            kind = IngresoTipo(nombre='Apertura importacion Miss Zapatos'); db.add(kind); db.flush()
        item_values = []
        for store_id, items in incoming.items():
            total = sum(cost*qty for _, _, qty, cost in items)
            entry = IngresoInventario(tipo_id=kind.id, bodega_id=store_id, fecha=day, moneda='CS', tasa_cambio=rate or None,
                                     observacion='Carga inicial Excel Miss Zapatos', usuario_registro=username,
                                     total_cs=total, total_usd=total/rate if rate else Decimal('0'))
            db.add(entry); db.flush()
            for pid, vid, qty, cost in items:
                usd = cost/rate if rate else Decimal('0')
                item_values.append(dict(ingreso_id=entry.id, producto_id=pid, variante_id=vid, cantidad=qty,
                                        costo_unitario_cs=cost, costo_unitario_usd=usd, subtotal_cs=cost*qty, subtotal_usd=usd*qty))
        _insert(db, IngresoItem, item_values)
    db.flush()
    return len(new_rows)
