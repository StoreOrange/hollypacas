"""Excel del Kardex: mismos movimientos y saldos que el reporte en pantalla."""
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def build_kardex_excel(start, end, warehouse, product, rows, summary):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Kardex'
    sheet.append(['Reporte Kardex por producto'])
    sheet.merge_cells('A1:N1')
    sheet['A1'].font = Font(size=16, bold=True, color='17365D')
    sheet.append(['Desde', start, 'Hasta', end])
    sheet['B2'].number_format = sheet['D2'].number_format = 'dd/mm/yyyy'
    sheet.append(['Bodega', warehouse])
    sheet.append(['Producto', product or 'Todos'])
    sheet.append(['Saldo inicial total', float(summary.get('saldo_inicial') or 0), 'Saldo final total', float(summary.get('saldo_final') or 0)])
    sheet['B5'].number_format = sheet['D5'].number_format = '#,##0.00'
    sheet.append(['Ingresos', summary.get('total_ingresos', 0), 'Egresos', summary.get('total_egresos', 0), 'Ventas', summary.get('total_ventas', 0)])
    sheet.append(['Los saldos de cada movimiento corresponden a su producto y bodega. Costos en córdobas (C$).'])
    sheet.merge_cells('A7:N7')
    sheet['A7'].font = Font(italic=True, color='595959')
    sheet.append(['Fecha', 'Tipo', 'Concepto', 'Sucursal', 'Bodega', 'Código', 'Producto', 'Saldo anterior', 'Entradas', 'Salidas', 'Saldo', 'Costo unitario C$', 'Costo total C$', 'Vendedor'])
    for row in rows:
        sheet.append([row.get('fecha'), row.get('tipo'), row.get('concepto'), row.get('branch'), row.get('bodega'), row.get('codigo'), row.get('descripcion'),
            *[float(row.get(key) or 0) for key in ('saldo_anterior', 'entrada', 'salida', 'saldo', 'costo_unit_cs', 'costo_total_cs')], row.get('vendedor') or '-'])
    # Códigos, conceptos y nombres son texto, incluso si comienzan por '='.
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = 's'
    for cell in sheet[8]:
        cell.fill = PatternFill('solid', fgColor='17365D')
        cell.font = Font(bold=True, color='FFFFFF')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    sheet.row_dimensions[8].height = 32
    for row in sheet.iter_rows(min_row=9):
        row[0].number_format = 'dd/mm/yyyy'
        for cell in row[7:13]:
            cell.number_format = '#,##0.00'
    for index, width in enumerate((13, 23, 38, 22, 24, 20, 48, 18, 16, 16, 18, 20, 20, 25), start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = 'H9'
    sheet.auto_filter.ref = f'A8:N{max(8, sheet.max_row)}'
    sheet.sheet_view.showGridLines = False
    sheet.print_title_rows = '1:8'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
