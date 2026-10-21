"""Kardex filters use isolated inventory; no real stock is modified."""
import json
import unittest
from datetime import date, datetime
from decimal import Decimal
from unittest.mock import patch
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from starlette.requests import Request
from fastapi import HTTPException
from app.database import Base
from app.routers import web as w
from app.main import app


class KardexTests(unittest.TestCase):
    def setUp(self):
        self.engine=create_engine('sqlite:///:memory:');Base.metadata.create_all(self.engine);self.db=Session(self.engine)
        self.company=patch.object(w,'get_active_company_key',return_value='bdzapatos');self.company.start()
        self.multi=patch.object(w,'_multi_branch_enabled_mode',return_value=True);self.multi.start()
        for i,code in [(1,'central'),(2,'kg'),(3,'kgf')]:
            self.db.add(w.Branch(id=i,code=code,name=code.upper()))
            self.db.add(w.Bodega(id=i,branch_id=i,code=code,name='Bodega '+code.upper()))
        self.db.add_all([w.Producto(id=1,cod_producto='ZAP-1',descripcion='Zapato casual'),w.Producto(id=2,cod_producto='ZAP-10',descripcion='Otro zapato'),w.ColorCatalog(id=1,nombre='Negro',abreviatura='NG')])
        for i,size in [(1,'38'),(2,'39')]:self.db.add(w.ShoeProductVariant(id=i,producto_id=1,color_id=1,talla=size,cod_variante='ZAP-1-'+size))
        self.db.add(w.IngresoTipo(id=1,nombre='Compra'))
        for i,warehouse,day,variant,qty in [(1,1,date(2026,10,1),1,5),(2,2,date(2026,10,1),1,20),(3,1,date(2026,10,19),1,2),(4,1,date(2026,10,19),2,9),(5,2,date(2026,10,19),1,3),(6,3,date(2026,10,19),1,4)]:
            self.db.add(w.IngresoInventario(id=i,bodega_id=warehouse,tipo_id=1,fecha=day,moneda='CS'))
            self.db.add(w.IngresoItem(ingreso_id=i,producto_id=1,variante_id=variant,cantidad=qty))
        self.db.commit()
        self.admin=w.User(email='admin',roles=[w.Role(name='administrador')],default_branch_id=1,default_bodega_id=1,branches=[self.db.get(w.Branch,1)])
        self.member=w.User(email='member',roles=[w.Role(name='contador')],default_branch_id=1,default_bodega_id=1,branches=[self.db.get(w.Branch,1)])

    def tearDown(self):self.db.close();self.engine.dispose();self.company.stop();self.multi.stop()
    def request(self,query=''):
        request=Request({'type':'http','method':'GET','path':'/reports/kardex','scheme':'http','server':('test',80),'headers':[], 'query_string':('start_date=2026-10-19&end_date=2026-10-19&'+query).encode(),'app':app,'router':app.router})
        request.state.branding={'company_key':'bdzapatos'}
        return request

    def test_admin_sees_three_warehouses_and_independent_balances(self):
        result=w._kardex_selection(self.request('variante_id=1'),self.db,self.admin)
        self.assertEqual([b.id for b in result[4]],[1,2,3])
        rows=result[5];self.assertEqual(len(rows),3)
        self.assertEqual({r['bodega_id']:r['saldo'] for r in rows},{1:Decimal(7),2:Decimal(23),3:Decimal(4)})
        self.assertEqual(rows[0]['codigo'],'ZAP-1-38')

    def test_all_does_not_leak_other_warehouses(self):
        rows=w._kardex_selection(self.request('variante_id=1'),self.db,self.member)[5]
        self.assertEqual(len(rows),1);self.assertEqual(rows[0]['saldo_anterior'],5)
        with self.assertRaises(HTTPException) as error:w._kardex_selection(self.request('bodega_id=2'),self.db,self.member)
        self.assertEqual(error.exception.status_code,403)

    def test_exact_product_and_variant_filters_do_not_match_partial_codes(self):
        result=w._kardex_selection(self.request('producto_id=2&producto=ZAP-1'),self.db,self.admin)
        self.assertEqual(result[5],[])
        result=w._kardex_selection(self.request('producto_id=1&bodega_id=1'),self.db,self.admin)
        self.assertEqual(len(result[5]),2)
        result=w._kardex_selection(self.request('variante_id=2&bodega_id=1'),self.db,self.admin)
        self.assertEqual(len(result[5]),1);self.assertEqual(result[5][0]['entrada'],9)
        self.assertEqual(result[5][0]['saldo_anterior'],0)

    def test_sales_and_egress_keep_variant_and_warehouse_scope(self):
        self.db.add(w.EgresoTipo(id=1,nombre='Salida'))
        for warehouse in [1,2]:
            self.db.add(w.VentaFactura(id=warehouse,numero=f'T-{warehouse}',bodega_id=warehouse,fecha=datetime(2026,10,19,10),moneda='CS',estado='ACTIVA'))
            self.db.add(w.EgresoInventario(id=warehouse,tipo_id=1,bodega_id=warehouse,fecha=date(2026,10,19),moneda='CS'))
            for variant in [1,2]:
                self.db.add(w.VentaItem(factura_id=warehouse,producto_id=1,variante_id=variant,cantidad=1))
                self.db.add(w.EgresoItem(egreso_id=warehouse,producto_id=1,variante_id=variant,cantidad=1))
        self.db.commit()
        rows=w._kardex_selection(self.request('variante_id=1'),self.db,self.member)[5]
        self.assertEqual(len(rows),3)
        self.assertEqual({r['bodega_id'] for r in rows},{1})
        self.assertEqual(sum(r['salida'] for r in rows),2)
        self.assertEqual(rows[-1]['saldo'],5)

    def test_smart_search_code_color_size_and_scannable_id(self):
        for q in ['ZAP-1-38','casual Negro 38','V000001']:
            data=json.loads(w.report_kardex_items(self.request(),q,self.db,self.admin).body)
            self.assertEqual([i['variante_id'] for i in data['items']],[1])
        data=json.loads(w.report_kardex_items(self.request(),'ZAP-1',self.db,self.admin).body)
        self.assertTrue(any(i['producto_id']==1 and i['variante_id'] is None for i in data['items']))

    def legacy_transfer(self, day=date(2026,10,19), typed=False):
        self.db.add(w.EgresoTipo(id=10,nombre='Traslado entre bodegas'))
        self.db.add(w.IngresoTipo(id=10,nombre='Traslado entre bodegas'))
        self.db.add(w.EgresoInventario(id=10,tipo_id=10,bodega_id=1,bodega_destino_id=2,fecha=day,moneda='CS'))
        self.db.add(w.IngresoInventario(id=10,tipo_id=10,bodega_id=2,fecha=day,moneda='CS',observacion='Traslado desde Central hacia KG. Egreso #10'))
        for variant,qty in [(1,2),(2,3)]:
            self.db.add(w.EgresoItem(egreso_id=10,producto_id=1,variante_id=variant,cantidad=qty,costo_unitario_cs=100))
            self.db.add(w.IngresoItem(ingreso_id=10,producto_id=1,variante_id=variant if typed else None,cantidad=qty,costo_unitario_cs=100))
        self.db.commit()

    def test_legacy_transfer_restores_only_selected_size_without_writes(self):
        self.legacy_transfer()
        rows,summary=w._build_kardex_movements(self.db,date(2026,10,19),date(2026,10,19),'2','',allowed_bodega_ids={2},producto_id=1,variante_id=1)
        self.assertEqual(summary['saldo_inicial'],20)
        self.assertEqual(summary['total_ingresos'],5)  # existing 3 + transferred 2
        self.assertEqual(summary['saldo_final'],25)
        recovered=[r for r in rows if r['source_id']==10]
        self.assertEqual(len(recovered),1);self.assertEqual(recovered[0]['entrada'],2)
        self.assertFalse(self.db.dirty);self.assertFalse(self.db.new)
        self.assertEqual(self.db.query(w.IngresoItem).filter_by(ingreso_id=10,variante_id=None).count(),2)
        # Scope also covers recovered rows; a forbidden destination stays absent.
        rows,_=w._build_kardex_movements(self.db,date(2026,10,19),date(2026,10,19),'all','',allowed_bodega_ids={1},producto_id=1,variante_id=1)
        self.assertEqual({r['bodega_id'] for r in rows},{1})

    def test_legacy_transfer_before_period_contributes_to_opening_balance(self):
        self.legacy_transfer(day=date(2026,10,18))
        rows,summary=w._build_kardex_movements(self.db,date(2026,10,19),date(2026,10,19),'2','',producto_id=1,variante_id=1)
        self.assertEqual(summary['saldo_inicial'],22);self.assertEqual(summary['saldo_final'],25)
        self.assertEqual(len(rows),1)

    def test_typed_transfer_is_not_reconstructed_twice(self):
        self.legacy_transfer(typed=True)
        _,summary=w._build_kardex_movements(self.db,date(2026,10,19),date(2026,10,19),'2','',producto_id=1,variante_id=1)
        self.assertEqual(summary['total_ingresos'],5)

    def test_partially_typed_transfer_only_recovers_missing_quantity(self):
        self.legacy_transfer()
        line=self.db.query(w.IngresoItem).filter_by(ingreso_id=10).order_by(w.IngresoItem.id).first()
        line.cantidad=1
        self.db.add(w.IngresoItem(ingreso_id=10,producto_id=1,variante_id=1,cantidad=1));self.db.commit()
        _,summary=w._build_kardex_movements(self.db,date(2026,10,19),date(2026,10,19),'2','',producto_id=1,variante_id=1)
        self.assertEqual(summary['total_ingresos'],5)

    def test_unproven_legacy_transfer_is_not_invented(self):
        from app.core.kardex_transfers import legacy_variant_receipts
        self.legacy_transfer()
        receipt=self.db.get(w.IngresoInventario,10)
        for bad_reference in ['Egreso #100','Egreso #10 | Egreso #11']:
            receipt.observacion=bad_reference;self.db.commit()
            self.assertEqual(legacy_variant_receipts(self.db,1,date(2026,10,19)),[])
        receipt.observacion='Egreso #10'
        self.db.query(w.IngresoItem).filter_by(ingreso_id=10).first().cantidad=99;self.db.commit()
        self.assertEqual(legacy_variant_receipts(self.db,1,date(2026,10,19)),[])

    def test_new_transfer_retains_variant_on_destination_receipt(self):
        import asyncio
        from starlette.datastructures import FormData
        self.db.add(w.EgresoTipo(id=10,nombre='Traslado entre bodegas'))
        self.db.add(w.ShoeVariantStock(variante_id=1,bodega_id=1,existencia=5));self.db.commit()
        request=self.request()
        request._form=FormData({'tipo_id':'10','bodega_id':'1','bodega_destino_id':'2','fecha':'2026-10-19','moneda':'CS','item_producto_id':'1','item_variante_id':'1','item_cantidad':'2','item_costo':'100'})
        with patch.object(w,'_build_auto_accounting_entry',return_value=None):
            response=asyncio.run(w.inventory_create_egreso(request,self.db,self.admin))
        self.assertNotIn('error=',response.headers['location'])
        destination=self.db.query(w.IngresoInventario).filter(w.IngresoInventario.observacion.like('%Egreso #%')).one()
        self.assertEqual(len(destination.items),1)
        self.assertEqual(destination.items[0].variante_id,1)
        self.assertEqual(destination.items[0].cantidad,2)
        self.assertEqual(self.db.query(w.ShoeVariantStock).filter_by(variante_id=1,bodega_id=2).one().existencia,2)

    def test_report_and_export_keep_the_same_filters(self):
        request=self.request('variante_id=1&bodega_id=2')
        response=w.report_kardex(request,self.db,self.admin)
        self.assertIn('kardex-item-modal',response.body.decode());self.assertIn('variante_id=1',response.body.decode())
        with patch.object(w,'_company_profile_payload',return_value={}):pdf=w.report_kardex_export(request,self.db,self.admin)
        self.assertEqual(pdf.media_type,'application/pdf')

    def excel_workbook(self, query='', user=None):
        import asyncio
        from io import BytesIO
        from openpyxl import load_workbook
        response=w.report_kardex_excel(self.request(query),self.db,user or self.admin)
        self.assertIn('attachment;',response.headers['content-disposition'])
        self.assertIn('.xlsx',response.headers['content-disposition'])
        async def read():
            return b''.join([chunk async for chunk in response.body_iterator])
        return load_workbook(BytesIO(asyncio.run(read())))

    def test_excel_matches_filtered_report_and_has_numeric_balances(self):
        query='variante_id=1&bodega_id=2'
        expected=w._kardex_selection(self.request(query),self.db,self.admin)
        sheet=self.excel_workbook(query).active
        self.assertEqual(sheet.max_row,9)
        self.assertEqual(sheet['B3'].value,'Bodega KG')
        self.assertIn('Talla 38',sheet['B4'].value)
        self.assertEqual(sheet['F9'].value,'ZAP-1-38')
        self.assertEqual(sheet['H9'].value,20)
        self.assertEqual(sheet['I9'].value,3)
        self.assertEqual(sheet['K9'].value,23)
        self.assertEqual(sheet['D5'].value,float(expected[6]['saldo_final']))
        self.assertEqual(sheet['K9'].data_type,'n')
        self.assertEqual(sheet['A9'].value.date(),date(2026,10,19))
        self.assertEqual(sheet.auto_filter.ref,'A8:N9')
        self.assertEqual(sheet.freeze_panes,'H9')

    def test_excel_empty_report_and_text_are_safe(self):
        sheet=self.excel_workbook('producto_id=2').active
        self.assertEqual(sheet.max_row,8)
        self.assertEqual(sheet['H8'].value,'Saldo anterior')
        self.db.get(w.Producto,1).descripcion='=1+1'
        self.db.commit()
        sheet=self.excel_workbook('producto_id=1&bodega_id=1').active
        self.assertEqual(sheet['G9'].value,'=1+1')
        self.assertEqual(sheet['G9'].data_type,'s')

    def test_excel_enforces_warehouse_scope_and_reports_permission(self):
        with patch.object(w,'_enforce_permission') as permission:
            sheet=self.excel_workbook('variante_id=1',self.member).active
            self.assertEqual(sheet.max_row,9)
            self.assertEqual(sheet['E9'].value,'Bodega CENTRAL')
            self.assertEqual(permission.call_args.args[2],'access.reports')
            with self.assertRaises(HTTPException) as error:
                w.report_kardex_excel(self.request('bodega_id=2'),self.db,self.member)
            self.assertEqual(error.exception.status_code,403)
