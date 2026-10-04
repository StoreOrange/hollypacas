import io
import unittest
from unittest.mock import patch
from starlette.requests import Request
from fastapi import UploadFile
from datetime import date
from decimal import Decimal
from openpyxl import Workbook,load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from app.database import Base
from app.routers import web
from app.core.shoe_inventory_import import HEADERS,parse,apply,template
from app.models.user import Branch
from app.models.inventory import Bodega,Producto,ShoeProductVariant,ShoeVariantStock,IngresoItem

class ImportTests(unittest.TestCase):
    def setUp(self):
        self.engine=create_engine('sqlite:///:memory:'); Base.metadata.create_all(self.engine)
        self.db=Session(self.engine)
        self.db.add(Branch(id=1,code='test',name='Test'))
        self.db.add_all([Bodega(id=i,branch_id=1,code=c,name=c.upper()) for i,c in [(1,'central'),(2,'kg'),(3,'kgf'),(4,'esteli')]])
        self.db.commit(); self.stores=self.db.query(Bodega).all()
    def tearDown(self):
        self.db.close(); self.engine.dispose()
    def content(self,rows,headers=None):
        w=Workbook();w.active.append(headers or HEADERS)
        for row in rows:w.active.append(row)
        s=io.BytesIO();w.save(s);return s.getvalue()
    def rows(self):
        return parse(self.content([
            ['#AIRCITBLU5.5','AIR 5.5 BLACK','CITY','5.5','BLACK PU','SANDALIA','#AIR',200,2,'NULL',None],
            ['#AIRCITBLU6','AIR 6 BLACK','CITY',6,'BLACK PU','SANDALIA','#AIR',300,None,1,'NULL']]))
    def test_loading_prices_warehouses_and_repeat(self):
        rows=self.rows()
        self.assertEqual(apply(self.db,rows,self.stores,Decimal('36'),date.today(),'test'),2)
        self.db.commit()
        products=self.db.query(Producto).order_by(Producto.id).all()
        self.assertEqual([p.precio_venta1 for p in products],[200,300])
        self.assertEqual(products[0].referencia_producto,'#AIR')
        variants=self.db.query(ShoeProductVariant).order_by(ShoeProductVariant.id).all()
        self.assertEqual(variants[0].talla,'5.5')
        stocks={(s.variante_id,s.bodega_id):s.existencia for s in self.db.query(ShoeVariantStock)}
        self.assertEqual(stocks[(variants[0].id,1)],2)
        self.assertEqual(stocks[(variants[1].id,3)],1)
        self.assertEqual(stocks[(variants[1].id,2)],0)
        self.assertFalse(any(b==4 for _,b in stocks))
        self.assertEqual(apply(self.db,rows,self.stores,Decimal('36'),date.today(),'test'),0)
        self.db.commit();self.assertEqual(self.db.query(IngresoItem).count(),2)
        rows[0]['001 - PRINCIPAL']=Decimal('3')
        with self.assertRaises(ValueError):apply(self.db,rows,self.stores,Decimal('36'),date.today(),'test')
        self.db.rollback();self.assertEqual(self.db.query(IngresoItem).count(),2)
    def test_cost_in_cordobas_and_blank_preserves(self):
        rows=self.rows();rows[0]['COSTO']=Decimal('464')
        apply(self.db,rows,self.stores,Decimal('36'),date.today(),'test');self.db.commit()
        self.assertEqual(self.db.query(Producto).first().costo_producto,464)
        rows[0]['COSTO']=None
        apply(self.db,rows,self.stores,Decimal('36'),date.today(),'test');self.db.commit()
        self.assertEqual(self.db.query(Producto).first().costo_producto,464)
    def test_invalid_and_duplicate_rows_rejected(self):
        raw=['SKU','Shoe','Brand',6,'Black','Line','BASE',200,1,0,0]
        for value in ['bad',-1,1.5,'NaN']:
            row=list(raw);row[8]=value
            with self.assertRaises(ValueError):parse(self.content([row]))
        with self.assertRaises(ValueError):parse(self.content([raw,raw]))
        with self.assertRaises(ValueError):parse(self.content([raw],HEADERS[:-1]))
    def test_route_isolation_and_preview_validation(self):
        request=Request({'type':'http','method':'GET','path':'/inventory/import/template','headers':[]})
        with patch.object(web,'_enforce_permission'), patch.object(web,'get_active_company_key',return_value='bdzapatos'):
            result=web.inventory_import_template(request,user=object())
            self.assertIn('inventario_miss_zapatos.xlsx',result.headers['content-disposition'])
            preview=web.inventory_import_preview(request,UploadFile(filename='test.xlsx',file=io.BytesIO(self.content([['bad']]))),user=object())
            self.assertEqual(preview.status_code,400)
        with patch.object(web,'_enforce_permission'), patch.object(web,'get_active_company_key',return_value='hollywood_pacas'):
            result=web.inventory_import_template(request,user=object())
            self.assertIn('productos_template.xlsx',result.headers['content-disposition'])

    def test_template_and_missing_warehouse(self):
        book=load_workbook(template());self.assertEqual([c.value for c in book.active[1]],HEADERS+['COSTO'])
        with self.assertRaises(ValueError):apply(self.db,self.rows(),self.stores[:2],Decimal('36'),date.today(),'test')
        self.assertEqual(self.db.query(Producto).count(),0)

if __name__=='__main__':unittest.main()
